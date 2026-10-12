import io
import re
import hashlib
import streamlit as st
import ollama

from pypdf import PdfReader
from pptx import Presentation
from docx import Document


# =========================================================
# CONFIG
# =========================================================

MODEL = "llama3.2"

st.set_page_config(
    page_title="AI Study Assistant",
    page_icon="🎓",
    layout="wide"
)


# =========================================================
# OLLAMA
# =========================================================

@st.cache_resource
def get_client():
    return ollama.Client()


client = get_client()


# =========================================================
# FILE EXTRACTION
# =========================================================

def extract_pdf(data):
    parts = []

    reader = PdfReader(io.BytesIO(data))

    for page in reader.pages:
        text = page.extract_text()

        if text:
            parts.append(text)

    return "\n".join(parts)


def extract_pptx(data):
    parts = []

    presentation = Presentation(io.BytesIO(data))

    for slide in presentation.slides:

        slide_parts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

        if slide_parts:
            parts.append("\n".join(slide_parts))

    return "\n\n".join(parts)


def extract_docx(data):
    parts = []

    document = Document(io.BytesIO(data))

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            parts.append(text)

    return "\n".join(parts)


def extract_text(data, filename):

    filename = filename.lower()

    if filename.endswith(".pdf"):
        return extract_pdf(data)

    if filename.endswith(".pptx"):
        return extract_pptx(data)

    if filename.endswith(".docx"):
        return extract_docx(data)

    return ""


# =========================================================
# TEXT
# =========================================================

def clean_text(text):

    text = re.sub(r"\s+", " ", text)

    return text.strip()


def create_chunks(text, chunk_size=1400, overlap=200):

    text = clean_text(text)

    if not text:
        return []

    chunks = []

    start = 0

    while start < len(text):

        end = start + chunk_size

        chunk = text[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start = end - overlap

    return chunks


# =========================================================
# RETRIEVAL
# =========================================================

def get_keywords(query):

    words = re.findall(
        r"[a-zA-Z0-9]+",
        query.lower()
    )

    stopwords = {
        "what", "is", "are", "the",
        "a", "an", "and", "or",
        "explain", "tell", "me",
        "about", "how", "why",
        "can", "you", "please",
        "give", "define", "of",
        "to", "in", "on", "for",
        "with", "from", "this",
        "that", "topic"
    }

    return [
        word
        for word in words
        if word not in stopwords
        and len(word) > 2
    ]


def retrieve_chunks(chunks, query, max_chunks=4):

    if not chunks:
        return ""

    keywords = get_keywords(query)

    if not keywords:
        return "\n\n".join(
            chunks[:max_chunks]
        )

    scored = []

    for chunk in chunks:

        lower = chunk.lower()

        score = 0

        for keyword in keywords:
            score += lower.count(keyword)

        scored.append(
            (score, chunk)
        )

    scored.sort(
        key=lambda x: x[0],
        reverse=True
    )

    selected = [
        chunk
        for score, chunk in scored[:max_chunks]
        if score > 0
    ]

    if not selected:
        selected = chunks[:max_chunks]

    return "\n\n".join(selected)


# =========================================================
# AI
# =========================================================

def stream_ai(prompt, max_tokens=280):

    stream = client.chat(

        model=MODEL,

        messages=[
            {
                "role": "system",
                "content": (
                    "You are a fast AI Study Assistant. "
                    "Answer clearly and directly. "
                    "Use simple B.Tech student-friendly "
                    "language. "
                    "Avoid unnecessary introductions."
                )
            },
            {
                "role": "user",
                "content": prompt
            }
        ],

        stream=True,

        options={
            "temperature": 0.2,
            "num_predict": max_tokens
        },

        keep_alive="15m"
    )

    for chunk in stream:

        content = chunk["message"]["content"]

        if content:
            yield content


def ai_text(prompt, max_tokens=300):

    response = client.chat(

        model=MODEL,

        messages=[
            {
                "role": "system",
                "content": (
                    "You are a concise AI Study Assistant. "
                    "Follow the requested format exactly. "
                    "Never output HTML."
                )
            },
            {
                "role": "user",
                "content": prompt
            }
        ],

        options={
            "temperature": 0.2,
            "num_predict": max_tokens
        },

        keep_alive="15m"
    )

    return response["message"]["content"]


# =========================================================
# POSTER HELPERS
# =========================================================

def extract_section(text, section, next_sections):

    pattern = (
        r"\b"
        + re.escape(section)
        + r"\s*:\s*"
        + r"(.*?)"
        + r"(?=\b(?:"
        + "|".join(
            re.escape(x)
            for x in next_sections
        )
        + r")\s*:|\Z)"
    )

    match = re.search(
        pattern,
        text,
        re.IGNORECASE | re.DOTALL
    )

    if match:
        return match.group(1).strip()

    return ""


def make_bullets(text):

    if not text:
        return []

    lines = text.splitlines()

    result = []

    for line in lines:

        line = line.strip()

        line = re.sub(
            r"^[-•*]\s*",
            "",
            line
        )

        if line:
            result.append(line)

    return result[:6]


def render_card(title, content, icon="📌"):

    with st.container(border=True):

        st.markdown(
            f"### {icon} {title}"
        )

        st.write(content)


def render_bullet_card(title, items, icon="📌"):

    with st.container(border=True):

        st.markdown(
            f"### {icon} {title}"
        )

        for item in items:

            st.markdown(
                f"• {item}"
            )


def render_poster(text, fallback_title):

    title = extract_section(
        text,
        "TITLE",
        [
            "DEFINITION",
            "KEY CONCEPTS",
            "IMPORTANT TERMS",
            "EXAMPLE",
            "APPLICATIONS",
            "EXAM TIP"
        ]
    )

    definition = extract_section(
        text,
        "DEFINITION",
        [
            "KEY CONCEPTS",
            "IMPORTANT TERMS",
            "EXAMPLE",
            "APPLICATIONS",
            "EXAM TIP"
        ]
    )

    concepts = extract_section(
        text,
        "KEY CONCEPTS",
        [
            "IMPORTANT TERMS",
            "EXAMPLE",
            "APPLICATIONS",
            "EXAM TIP"
        ]
    )

    terms = extract_section(
        text,
        "IMPORTANT TERMS",
        [
            "EXAMPLE",
            "APPLICATIONS",
            "EXAM TIP"
        ]
    )

    example = extract_section(
        text,
        "EXAMPLE",
        [
            "APPLICATIONS",
            "EXAM TIP"
        ]
    )

    applications = extract_section(
        text,
        "APPLICATIONS",
        [
            "EXAM TIP"
        ]
    )

    exam_tip = extract_section(
        text,
        "EXAM TIP",
        []
    )

    if not title:
        title = fallback_title

    if not definition:
        definition = "Key concept from the selected study material."

    if not example:
        example = "Use the concepts above to understand the topic."

    concept_items = make_bullets(concepts)

    term_items = make_bullets(terms)

    application_items = make_bullets(applications)

    # -----------------------------------------------------
    # POSTER HEADER
    # -----------------------------------------------------

    st.markdown(
        f"""
        # 🎓 {title}

        **AI STUDY • QUICK REVISION • EXAM READY**
        """
    )

    st.divider()

    # -----------------------------------------------------
    # DEFINITION
    # -----------------------------------------------------

    render_card(
        "What Is It?",
        definition,
        "💡"
    )

    # -----------------------------------------------------
    # CONCEPTS + TERMS
    # -----------------------------------------------------

    col1, col2 = st.columns(2)

    with col1:

        if concept_items:

            render_bullet_card(
                "Key Concepts",
                concept_items,
                "🧠"
            )

        else:

            render_card(
                "Key Concepts",
                "No separate key concepts were generated.",
                "🧠"
            )

    with col2:

        if term_items:

            render_bullet_card(
                "Important Terms",
                term_items,
                "🔑"
            )

        else:

            render_card(
                "Important Terms",
                "No separate terms were generated.",
                "🔑"
            )

    # -----------------------------------------------------
    # EXAMPLE + APPLICATIONS
    # -----------------------------------------------------

    col1, col2 = st.columns(2)

    with col1:

        render_card(
            "Example",
            example,
            "🌍"
        )

    with col2:

        if application_items:

            render_bullet_card(
                "Applications",
                application_items,
                "🚀"
            )

        else:

            render_card(
                "Applications",
                "Applications were not specified.",
                "🚀"
            )

    # -----------------------------------------------------
    # EXAM TIP
    # -----------------------------------------------------

    st.divider()

    if exam_tip:

        st.success(
            f"🎯 **Exam Tip:** {exam_tip}"
        )

    else:

        st.info(
            "🎯 **Exam Tip:** Revise the definition, "
            "key concepts and important terms."
        )


# =========================================================
# FLASHCARDS
# =========================================================

def render_flashcards(text):

    pattern = (
        r"CARD\s*(\d+)"
        r"\s*Q:\s*(.*?)"
        r"\s*A:\s*(.*?)"
        r"(?=CARD\s*\d+\s*Q:|\Z)"
    )

    matches = re.findall(
        pattern,
        text,
        re.IGNORECASE | re.DOTALL
    )

    if not matches:

        st.write(text)
        return

    columns = st.columns(2)

    for index, (
        number,
        question,
        answer
    ) in enumerate(matches):

        with columns[index % 2]:

            with st.container(border=True):

                st.markdown(
                    f"### 🃏 Flashcard {number}"
                )

                st.markdown(
                    f"**Q:** {question.strip()}"
                )

                st.markdown(
                    f"**A:** {answer.strip()}"
                )


# =========================================================
# STUDY PLAN
# =========================================================

def render_study_plan(text):

    pattern = (
        r"DAY\s*(\d+)"
        r"\s*:\s*TOPIC\s*:\s*(.*?)"
        r"\s*TASK\s*:\s*(.*?)"
        r"(?=DAY\s*\d+\s*:|\Z)"
    )

    matches = re.findall(
        pattern,
        text,
        re.IGNORECASE | re.DOTALL
    )

    if not matches:

        st.write(text)
        return

    for day, topic, task in matches:

        with st.container(border=True):

            st.markdown(
                f"### 📅 Day {day}"
            )

            st.markdown(
                f"**📚 Topic:** {topic.strip()}"
            )

            st.markdown(
                f"**🎯 Task:** {task.strip()}"
            )


# =========================================================
# HEADER
# =========================================================

st.title(
    "🎓 AI Study Assistant"
)

st.caption(
    "Learn smarter from your study material or any topic."
)


# =========================================================
# MODE
# =========================================================

mode = st.radio(
    "Choose Study Mode",

    [
        "📚 Study Material Available",
        "💡 Topic Only"
    ],

    horizontal=True
)


# =========================================================
# STUDY MATERIAL
# =========================================================

if mode == "📚 Study Material Available":

    st.subheader(
        "📂 Upload Study Material"
    )

    uploaded_file = st.file_uploader(
        "Upload PDF, PPTX or DOCX",
        type=[
            "pdf",
            "pptx",
            "docx"
        ]
    )

    if uploaded_file:

        file_bytes = uploaded_file.getvalue()

        file_hash = hashlib.md5(
            file_bytes
        ).hexdigest()

        if (
            st.session_state.get(
                "file_hash"
            )
            != file_hash
        ):

            with st.spinner(
                "⚡ Reading study material..."
            ):

                text = extract_text(
                    file_bytes,
                    uploaded_file.name
                )

                text = clean_text(text)

                chunks = create_chunks(text)

                st.session_state.file_hash = file_hash

                st.session_state.document_text = text

                st.session_state.chunks = chunks

        chunks = st.session_state.get(
            "chunks",
            []
        )

        st.success(
            f"✅ {uploaded_file.name} loaded"
        )

        st.caption(
            f"📄 {len(chunks)} searchable sections"
        )

        question = st.text_input(
            "Enter your question or topic",
            placeholder="Example: Explain Critical Section"
        )

        action = st.selectbox(
            "Choose Action",

            [
                "Ask Question",
                "Generate Summary",
                "Generate Short Notes",
                "Generate MCQs",
                "Explain Difficult Topic",
                "🃏 Generate Flashcards",
                "📅 Create Study Plan",
                "🖼️ Create Visual Poster"
            ]
        )

        if st.button(
            "🚀 Generate",
            type="primary",
            use_container_width=True
        ):

            if not chunks:

                st.error(
                    "Could not read this file."
                )

            elif action == "Ask Question":

                if not question.strip():

                    st.warning(
                        "Enter a question first."
                    )

                else:

                    context = retrieve_chunks(
                        chunks,
                        question,
                        4
                    )

                    prompt = f"""
Answer using ONLY this uploaded study material.

MATERIAL:
{context}

QUESTION:
{question}

If the answer is not present, say:
"This information is not available in the uploaded study material."

Keep the answer concise.
"""

                    st.subheader(
                        "🤖 Answer"
                    )

                    st.write_stream(
                        stream_ai(
                            prompt,
                            250
                        )
                    )

            elif action == "Generate Summary":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "main topics",
                    5
                )

                prompt = f"""
Create a concise exam-friendly summary.

Use ONLY:
{context}

Include definitions and key points.
"""

                st.subheader(
                    "📝 Summary"
                )

                st.write_stream(
                    stream_ai(
                        prompt,
                        300
                    )
                )

            elif action == "Generate Short Notes":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "important topics",
                    5
                )

                prompt = f"""
Create short revision notes using ONLY:

{context}

Include:
Definition
Key points
Important terms
Exam points
"""

                st.subheader(
                    "📌 Short Notes"
                )

                st.write_stream(
                    stream_ai(
                        prompt,
                        350
                    )
                )

            elif action == "Generate MCQs":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "important concepts",
                    5
                )

                prompt = f"""
Create 5 MCQs using ONLY:

{context}

For each:
A)
B)
C)
D)
Answer:
"""

                st.subheader(
                    "❓ MCQs"
                )

                st.write_stream(
                    stream_ai(
                        prompt,
                        450
                    )
                )

            elif action == "Explain Difficult Topic":

                if not question.strip():

                    st.warning(
                        "Enter a topic first."
                    )

                else:

                    context = retrieve_chunks(
                        chunks,
                        question,
                        4
                    )

                    prompt = f"""
Explain the topic using ONLY:

{context}

TOPIC:
{question}

Give:
1. Simple definition
2. Main idea
3. Important points
4. Example if available

Keep it short.
"""

                    st.subheader(
                        "📖 Explanation"
                    )

                    st.write_stream(
                        stream_ai(
                            prompt,
                            300
                        )
                    )

            elif action == "🃏 Generate Flashcards":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "important concepts",
                    5
                )

                prompt = f"""
Create 6 flashcards using ONLY:

{context}

Format exactly:

CARD 1
Q: ...
A: ...

CARD 2
Q: ...
A: ...

CARD 3
Q: ...
A: ...

CARD 4
Q: ...
A: ...

CARD 5
Q: ...
A: ...

CARD 6
Q: ...
A: ...
"""

                result = ai_text(
                    prompt,
                    450
                )

                st.subheader(
                    "🃏 Flashcards"
                )

                render_flashcards(
                    result
                )

            elif action == "📅 Create Study Plan":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "all important topics",
                    6
                )

                prompt = f"""
Create a 5-day study plan using ONLY:

{context}

Format:

DAY 1:
TOPIC:
...
TASK:
...

DAY 2:
TOPIC:
...
TASK:
...

DAY 3:
TOPIC:
...
TASK:
...

DAY 4:
TOPIC:
...
TASK:
...

DAY 5:
TOPIC:
...
TASK:
...
"""

                result = ai_text(
                    prompt,
                    400
                )

                st.subheader(
                    "📅 Study Plan"
                )

                render_study_plan(
                    result
                )

            elif action == "🖼️ Create Visual Poster":

                context = retrieve_chunks(
                    chunks,
                    question
                    if question.strip()
                    else "main important concepts",
                    5
                )

                prompt = f"""
Create content for a visual study poster.

Use ONLY this material:

{context}

Use EXACTLY:

TITLE:
...

DEFINITION:
...

KEY CONCEPTS:
- ...
- ...
- ...
- ...

IMPORTANT TERMS:
- ...
- ...
- ...
- ...

EXAMPLE:
...

APPLICATIONS:
- ...
- ...
- ...

EXAM TIP:
...

Keep everything concise.
Do NOT output HTML.
"""

                result = ai_text(
                    prompt,
                    350
                )

                st.subheader(
                    "🖼️ Visual Study Poster"
                )

                render_poster(
                    result,
                    question
                    if question.strip()
                    else "Study Material"
                )


# =========================================================
# TOPIC ONLY
# =========================================================

else:

    st.subheader(
        "💡 Learn Any Topic"
    )

    topic = st.text_input(
        "Enter Topic Name",
        placeholder="Example: Knowledge-Based Agents"
    )

    action = st.selectbox(
        "Choose Action",

        [
            "Learn Topic",
            "Generate Short Notes",
            "Generate MCQs",
            "Important Exam Questions",
            "Explain Like a Beginner",
            "📅 Create Study Plan",
            "🃏 Generate Flashcards",
            "🖼️ Create Visual Poster"
        ]
    )

    if st.button(
        "⚡ Generate",
        type="primary",
        use_container_width=True
    ):

        if not topic.strip():

            st.warning(
                "Enter a topic first."
            )

        elif action == "Learn Topic":

            prompt = f"""
Teach me about {topic}.

Give:
1. Simple definition
2. Main concept
3. 4 important points
4. One example
5. One exam point

Keep it concise.
"""

            st.subheader(
                "📚 Learn Topic"
            )

            st.write_stream(
                stream_ai(
                    prompt,
                    300
                )
            )

        elif action == "Generate Short Notes":

            prompt = f"""
Create short exam revision notes for:

{topic}

Include:
Definition
Key concepts
Important points
Exam keywords
"""

            st.subheader(
                "📌 Short Notes"
            )

            st.write_stream(
                stream_ai(
                    prompt,
                    300
                )
            )

        elif action == "Generate MCQs":

            prompt = f"""
Create 5 MCQs about:

{topic}

For each:
A)
B)
C)
D)
Answer:
"""

            st.subheader(
                "❓ MCQs"
            )

            st.write_stream(
                stream_ai(
                    prompt,
                    400
                )
            )

        elif action == "Important Exam Questions":

            prompt = f"""
Give 5 important exam questions about:

{topic}

Include:
2 short-answer
2 conceptual
1 long-answer

Do not explain.
"""

            st.subheader(
                "🎯 Important Exam Questions"
            )

            st.write_stream(
                stream_ai(
                    prompt,
                    250
                )
            )

        elif action == "Explain Like a Beginner":

            prompt = f"""
Explain {topic} like a complete beginner.

Use:
Simple language
One analogy
3 important points
One example

Keep it short.
"""

            st.subheader(
                "🧠 Beginner Explanation"
            )

            st.write_stream(
                stream_ai(
                    prompt,
                    280
                )
            )

        elif action == "📅 Create Study Plan":

            prompt = f"""
Create a 5-day study plan for:

{topic}

Format:

DAY 1:
TOPIC:
...
TASK:
...

DAY 2:
TOPIC:
...
TASK:
...

DAY 3:
TOPIC:
...
TASK:
...

DAY 4:
TOPIC:
...
TASK:
...

DAY 5:
TOPIC:
...
TASK:
...
"""

            result = ai_text(
                prompt,
                400
            )

            st.subheader(
                "📅 Study Plan"
            )

            render_study_plan(
                result
            )

        elif action == "🃏 Generate Flashcards":

            prompt = f"""
Create 6 flashcards about:

{topic}

Format:

CARD 1
Q: ...
A: ...

CARD 2
Q: ...
A: ...

CARD 3
Q: ...
A: ...

CARD 4
Q: ...
A: ...

CARD 5
Q: ...
A: ...

CARD 6
Q: ...
A: ...
"""

            result = ai_text(
                prompt,
                450
            )

            st.subheader(
                "🃏 Flashcards"
            )

            render_flashcards(
                result
            )

        elif action == "🖼️ Create Visual Poster":

            prompt = f"""
Create content for a visual study poster about:

{topic}

Use EXACTLY:

TITLE:
...

DEFINITION:
...

KEY CONCEPTS:
- ...
- ...
- ...
- ...

IMPORTANT TERMS:
- ...
- ...
- ...
- ...

EXAMPLE:
...

APPLICATIONS:
- ...
- ...
- ...

EXAM TIP:
...

Keep everything concise.
Do NOT output HTML.
"""

            result = ai_text(
                prompt,
                350
            )

            st.subheader(
                "🖼️ Visual Study Poster"
            )

            render_poster(
                result,
                topic
            )


# =========================================================
# FOOTER
# =========================================================

st.divider()

st.caption(
    "⚡ Powered by Ollama + Llama 3.2 | "
    "Local AI Study Assistant"
)
